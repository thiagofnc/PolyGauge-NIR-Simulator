"""Tests for the shared pipeline, loaders, Excel export and UI/core/export consistency."""

import math
import os
import tempfile
import unittest

import numpy as np

from BroadbandExcel import export_broadband_workbook
from BroadbandPipeline import (BASELINE_AUTO, BASELINE_OFF, BASELINE_ON, DETECTOR_RESPONSIVITY, LAYERS_PHYSICAL,
                               LAYERS_RATIO, REFERENCE_MULTIPLIER, SOURCE_BLACKBODY, BroadbandConfig,
                               run_broadband_analysis)
from SpectralData import load_material_library, load_source_presets, load_spectral_curve_csv, load_material_spectrum

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))


def synthetic_library(absorbance=0.2, low=2000.0, high=3000.0, reference_um=None):
    wl = np.linspace(low, high, 201)
    values = absorbance(wl) if callable(absorbance) else np.full_like(wl, absorbance)
    return {"Film": {"name": "Film", "wavelength_nm": wl, "absorbance": values, "convention": "base10",
                     "reference_thickness_um": reference_um, "source": "synthetic", "note": ""}}


def config(**overrides):
    base = dict(material="Film", values=[1, 2, 3, 4], no_film_voltage_mv=100.0, layer_to_reference_ratio=1.0,
                wavelength_min_nm=2000.0, wavelength_max_nm=3000.0)
    base.update(overrides)
    return BroadbandConfig(**base)


class PipelineTests(unittest.TestCase):
    def test_clean_data_is_left_completely_alone_by_default(self):
        cfg = config()
        self.assertEqual(cfg.baseline_correction, BASELINE_AUTO)
        self.assertFalse(cfg.clamp_negative or cfg.interface_correction)
        self.assertEqual(cfg.dark_voltage_mv, 0.0)
        result = run_broadband_analysis(cfg, synthetic_library())
        self.assertFalse(result["corrections"]["any_enabled"])
        self.assertEqual(result["corrections"]["baseline_offset"], 0.0)
        for row in result["rows"]:
            self.assertEqual(row["predicted_voltage_mv"], row["raw_predicted_voltage_mv"])

    def test_physical_thickness_uses_x_over_x_ref(self):
        result = run_broadband_analysis(config(thickness_mode=LAYERS_PHYSICAL, layer_thickness_um=25.0),
                                        synthetic_library(0.2, reference_um=50.0))
        row = result["rows"][3]  # 4 layers x 25 µm = 100 µm = 2 x_ref
        self.assertEqual(row["thickness_um"], 100.0)
        self.assertEqual(row["thickness_scale"], 2.0)
        self.assertAlmostEqual(row["predicted_transmission"], 10 ** (-2 * 0.2), places=12)
        self.assertTrue(result["thickness"]["reports_um"])

    def test_layer_thickness_value_changes_prediction(self):
        library = synthetic_library(0.2, reference_um=50.0)
        thin = run_broadband_analysis(config(thickness_mode=LAYERS_PHYSICAL, layer_thickness_um=10.0), library)
        thick = run_broadband_analysis(config(thickness_mode=LAYERS_PHYSICAL, layer_thickness_um=40.0), library)
        self.assertGreater(thin["rows"][0]["predicted_voltage_mv"], thick["rows"][0]["predicted_voltage_mv"])

    def test_physical_mode_refuses_unknown_reference_thickness(self):
        with self.assertRaises(ValueError):
            run_broadband_analysis(config(thickness_mode=LAYERS_PHYSICAL, layer_thickness_um=25.0), synthetic_library())

    def test_ratio_mode_never_reports_micrometres(self):
        result = run_broadband_analysis(config(layer_to_reference_ratio=0.5), synthetic_library(0.2))
        self.assertFalse(result["thickness"]["reports_um"])
        self.assertIsNone(result["rows"][0]["thickness_um"])
        self.assertAlmostEqual(result["rows"][3]["predicted_transmission"], 10 ** (-0.4), places=12)

    def test_dark_voltage_in_pipeline(self):
        result = run_broadband_analysis(config(dark_voltage_mv=12.0, measured_voltages_by_layer={1: 60.0}),
                                        synthetic_library(0.2))
        row = result["rows"][0]
        self.assertAlmostEqual(row["predicted_voltage_mv"], 12.0 + 88.0 * 10 ** -0.2, places=12)
        self.assertAlmostEqual(row["measured_transmission"], (60.0 - 12.0) / 88.0, places=12)

    def test_measured_layers_not_matched_in_multiplier_mode(self):
        result = run_broadband_analysis(config(thickness_mode=REFERENCE_MULTIPLIER,
                                               measured_voltages_by_layer={1: 50.0, 2: 40.0}), synthetic_library())
        self.assertTrue(all(row["measured_voltage_mv"] is None for row in result["rows"]))
        self.assertEqual(result["metrics"]["n"], 0)
        self.assertTrue(any("NOT compared" in w for w in result["warnings"]))

    def test_measured_comparison_and_metrics(self):
        library = synthetic_library(0.2)
        result = run_broadband_analysis(config(values=[1, 2], measured_voltages_by_layer={1: 60.0, 2: 40.0}), library)
        predicted = [10 ** -0.2 * 100, 10 ** -0.4 * 100]
        errors = [predicted[0] - 60.0, predicted[1] - 40.0]
        self.assertAlmostEqual(result["rows"][0]["error_mv"], errors[0], places=10)
        self.assertAlmostEqual(result["rows"][1]["percent_error"], 100 * errors[1] / 40.0, places=10)
        self.assertAlmostEqual(result["metrics"]["mae_mv"], np.mean(np.abs(errors)), places=10)
        self.assertAlmostEqual(result["metrics"]["rmse_mv"], np.sqrt(np.mean(np.square(errors))), places=10)

    def test_inspected_value_is_never_snapped(self):
        with self.assertRaises(ValueError):
            run_broadband_analysis(config(inspect_value=2.2), synthetic_library())
        result = run_broadband_analysis(config(inspect_value=3), synthetic_library())
        self.assertEqual(result["inspected_index"], 2)

    def test_partial_coverage_flagged_and_reported(self):
        result = run_broadband_analysis(config(wavelength_max_nm=4000.0), synthetic_library(0.2))
        self.assertTrue(result["coverage"]["partial"])
        self.assertAlmostEqual(result["coverage"]["fraction"], 0.5, places=12)
        self.assertTrue(result["warnings"][0].startswith("PARTIAL SPECTRAL COVERAGE"))
        low, high = result["rows"][0]["predicted_voltage_bounds_mv"]
        self.assertLess(low, result["rows"][0]["predicted_voltage_mv"])
        self.assertGreater(high, result["rows"][0]["predicted_voltage_mv"])

    def test_negative_absorbance_warning_only_counts_active_region(self):
        library = synthetic_library(lambda wl: np.where(wl < 2500, -0.05, 0.1))
        inside = run_broadband_analysis(config(baseline_correction=BASELINE_OFF), library)
        outside = run_broadband_analysis(config(wavelength_min_nm=2600.0, baseline_correction=BASELINE_OFF), library)
        self.assertGreater(inside["negative_absorbance"]["points"], 0)
        self.assertEqual(outside["negative_absorbance"]["points"], 0)

    def test_enabled_corrections_are_reported(self):
        library = synthetic_library(lambda wl: np.where(wl < 2500, -0.05, 0.1))
        result = run_broadband_analysis(config(clamp_negative=True, interface_correction=True, refractive_index=1.5,
                                               baseline_correction=BASELINE_OFF), library)
        corrections = result["corrections"]
        self.assertTrue(corrections["any_enabled"])
        self.assertGreater(corrections["clamped_points"], 0)
        self.assertAlmostEqual(corrections["interface_per_layer"], 0.96 ** 2)
        row = result["rows"][1]
        self.assertAlmostEqual(row["interface_transmission"], 0.96 ** 4, places=12)
        self.assertNotEqual(row["predicted_voltage_mv"], row["raw_predicted_voltage_mv"])

    def test_interface_correction_needs_layer_mode(self):
        with self.assertRaises(ValueError):
            run_broadband_analysis(config(thickness_mode=REFERENCE_MULTIPLIER, interface_correction=True,
                                          refractive_index=1.5), synthetic_library())

    def test_responsivity_curve_zero_outside_range(self):
        response = {"wavelength_nm": np.array([2200.0, 2800.0]), "values": np.array([1.0, 1.0]), "note": "test"}
        library = synthetic_library(lambda wl: np.where(wl < 2200, 5.0, 0.1))
        result = run_broadband_analysis(config(detector_mode=DETECTOR_RESPONSIVITY, detector_responsivity=response,
                                               wavelength_min_nm=None, wavelength_max_nm=None), library)
        self.assertAlmostEqual(result["rows"][0]["spectral_transmission"], 10 ** -0.1, places=10)
        self.assertEqual(result["weighting"]["weight_domain_nm"], (2200.0, 2800.0))

    def test_blackbody_source_changes_prediction(self):
        library = synthetic_library(lambda wl: (wl - 2000) / 1000)
        flat = run_broadband_analysis(config(), library)
        hot = run_broadband_analysis(config(source_mode=SOURCE_BLACKBODY, source_temperature_k=903.15), library)
        self.assertNotAlmostEqual(flat["rows"][0]["spectral_transmission"], hot["rows"][0]["spectral_transmission"], 4)
        with self.assertRaises(ValueError):
            run_broadband_analysis(config(source_mode=SOURCE_BLACKBODY), library)


class AutomaticBaselineTests(unittest.TestCase):
    """The baseline is corrected when the data needs it, never silently otherwise."""

    def negative_baseline_library(self):
        # Transparent everywhere except one band, with the transparent part sitting below zero.
        return synthetic_library(lambda wl: np.where((wl > 2400) & (wl < 2600), 0.4, -0.05))

    def test_auto_subtracts_a_negative_baseline(self):
        result = run_broadband_analysis(config(), self.negative_baseline_library())
        corrections = result["corrections"]
        self.assertTrue(corrections["baseline_enabled"])
        self.assertTrue(corrections["baseline_automatic"])
        self.assertAlmostEqual(corrections["baseline_offset"], -0.05, places=6)
        self.assertIn("below zero", corrections["baseline_reason"])
        self.assertTrue(any(w.startswith("BASELINE CORRECTED AUTOMATICALLY") for w in result["warnings"]))

    def test_auto_correction_restores_a_falling_voltage_curve(self):
        library = self.negative_baseline_library()
        raw = run_broadband_analysis(config(baseline_correction=BASELINE_OFF), library)
        fixed = run_broadband_analysis(config(), library)
        raw_voltages = [row["predicted_voltage_mv"] for row in raw["rows"]]
        fixed_voltages = [row["predicted_voltage_mv"] for row in fixed["rows"]]
        self.assertLess(raw_voltages[0], raw_voltages[-1])      # unusable: brighter through more film
        self.assertGreater(fixed_voltages[0], fixed_voltages[-1])
        self.assertTrue(all(b < a for a, b in zip(fixed_voltages, fixed_voltages[1:])))

    def test_auto_never_removes_real_absorption(self):
        # An all-positive spectrum has a positive baseline; subtracting it would delete absorption.
        result = run_broadband_analysis(config(), synthetic_library(0.2))
        self.assertFalse(result["corrections"]["baseline_enabled"])
        self.assertEqual(result["corrections"]["baseline_offset"], 0.0)
        self.assertIn("could remove real absorption", result["corrections"]["baseline_reason"])
        forced = run_broadband_analysis(config(baseline_correction=BASELINE_ON), synthetic_library(0.2))
        self.assertTrue(forced["corrections"]["baseline_enabled"])
        self.assertAlmostEqual(forced["rows"][0]["predicted_transmission"], 1.0, places=9)

    def test_auto_ignores_baseline_noise_within_tolerance(self):
        result = run_broadband_analysis(config(), synthetic_library(lambda wl: np.where(wl < 2500, -0.0005, 0.3)))
        self.assertFalse(result["corrections"]["baseline_enabled"])
        self.assertIn("within noise of zero", result["corrections"]["baseline_reason"])

    def test_off_and_on_settings_are_honoured(self):
        library = self.negative_baseline_library()
        off = run_broadband_analysis(config(baseline_correction=BASELINE_OFF), library)
        self.assertFalse(off["corrections"]["baseline_enabled"])
        self.assertFalse(off["corrections"]["baseline_automatic"])
        on = run_broadband_analysis(config(baseline_correction=BASELINE_ON), library)
        self.assertTrue(on["corrections"]["baseline_enabled"])
        self.assertFalse(on["corrections"]["baseline_automatic"])
        self.assertFalse(run_broadband_analysis(config(baseline_correction=False), library)["corrections"]["baseline_enabled"])
        self.assertTrue(run_broadband_analysis(config(baseline_correction=True), library)["corrections"]["baseline_enabled"])
        with self.assertRaises(ValueError):
            run_broadband_analysis(config(baseline_correction="sometimes"), library)

    def test_impossible_rising_prediction_is_called_out(self):
        result = run_broadband_analysis(config(baseline_correction=BASELINE_OFF), self.negative_baseline_library())
        self.assertTrue(result["warnings"][0].startswith("PREDICTED VOLTAGE RISES WITH THICKNESS"))
        fixed = run_broadband_analysis(config(), self.negative_baseline_library())
        self.assertFalse(any(w.startswith("PREDICTED VOLTAGE RISES") for w in fixed["warnings"]))

    def test_bundled_pe_spectrum_is_corrected_and_falls(self):
        library = load_material_library(BASE_DIR, include_discovered=False)
        cfg = BroadbandConfig(material="PE", values=[1, 2, 3, 4], no_film_voltage_mv=498.0,
                              layer_to_reference_ratio=1.0, wavelength_min_nm=2000.0, wavelength_max_nm=12000.0)
        result = run_broadband_analysis(cfg, library)
        self.assertTrue(result["corrections"]["baseline_automatic"])
        voltages = [row["predicted_voltage_mv"] for row in result["rows"]]
        self.assertTrue(all(b < a for a, b in zip(voltages, voltages[1:])), voltages)
        self.assertTrue(all(v < cfg.no_film_voltage_mv for v in voltages))


class ThicknessEstimateTests(unittest.TestCase):
    """Inverse mode must be the forward model solved backwards, with the same settings."""

    def test_estimate_inverts_the_forward_prediction(self):
        library = synthetic_library(0.2)
        forward = run_broadband_analysis(config(), library)
        target = forward["rows"][2]["predicted_voltage_mv"]  # 3 layers
        result = run_broadband_analysis(config(target_voltage_mv=target), library)
        estimate = result["estimate"]
        self.assertEqual(estimate["status"], "ok")
        self.assertAlmostEqual(estimate["thickness_scale"], 3.0, places=6)
        self.assertAlmostEqual(estimate["layer_equivalent"], 3.0, places=6)
        self.assertAlmostEqual(estimate["back_predicted_voltage_mv"], target, places=6)
        self.assertLess(abs(estimate["residual_mv"]), 1e-6)

    def test_forward_rows_are_unaffected_by_the_estimate(self):
        library = synthetic_library(0.2)
        plain = run_broadband_analysis(config(), library)
        with_estimate = run_broadband_analysis(config(target_voltage_mv=50.0), library)
        self.assertIsNone(plain["estimate"])
        for a, b in zip(plain["rows"], with_estimate["rows"]):
            self.assertEqual(a["predicted_voltage_mv"], b["predicted_voltage_mv"])

    def test_estimate_reports_micrometres_only_with_a_real_reference_thickness(self):
        known = run_broadband_analysis(config(thickness_mode=LAYERS_PHYSICAL, layer_thickness_um=25.0,
                                              target_voltage_mv=10 ** -0.4 * 100),
                                       synthetic_library(0.2, reference_um=50.0))["estimate"]
        self.assertAlmostEqual(known["thickness_scale"], 2.0, places=6)
        self.assertAlmostEqual(known["thickness_um"], 100.0, places=4)
        self.assertAlmostEqual(known["layer_equivalent"], 4.0, places=4)  # 25 µm layers
        unknown = run_broadband_analysis(config(target_voltage_mv=10 ** -0.4 * 100),
                                         synthetic_library(0.2))["estimate"]
        self.assertIsNone(unknown["thickness_um"])
        self.assertAlmostEqual(unknown["thickness_scale"], 2.0, places=6)

    def test_dark_voltage_is_used_in_the_inverse(self):
        library = synthetic_library(0.2)
        forward = run_broadband_analysis(config(dark_voltage_mv=12.0), library)
        target = forward["rows"][1]["predicted_voltage_mv"]
        estimate = run_broadband_analysis(config(dark_voltage_mv=12.0, target_voltage_mv=target), library)["estimate"]
        self.assertAlmostEqual(estimate["thickness_scale"], 2.0, places=6)
        ignoring_dark = run_broadband_analysis(config(target_voltage_mv=target), library)["estimate"]
        self.assertNotAlmostEqual(ignoring_dark["thickness_scale"], 2.0, places=3)

    def test_interface_correction_is_applied_in_the_inverse(self):
        library = synthetic_library(0.2)
        settings = dict(interface_correction=True, refractive_index=1.5)
        forward = run_broadband_analysis(config(**settings), library)
        target = forward["rows"][2]["predicted_voltage_mv"]
        estimate = run_broadband_analysis(config(target_voltage_mv=target, **settings), library)["estimate"]
        self.assertAlmostEqual(estimate["thickness_scale"], 3.0, places=5)
        self.assertLess(estimate["interface_transmission"], 1.0)

    def test_source_weighting_is_applied_in_the_inverse(self):
        library = synthetic_library(lambda wl: (wl - 2000) / 1000)
        settings = dict(source_mode=SOURCE_BLACKBODY, source_temperature_k=903.15)
        target = run_broadband_analysis(config(**settings), library)["rows"][1]["predicted_voltage_mv"]
        weighted = run_broadband_analysis(config(target_voltage_mv=target, **settings), library)["estimate"]
        flat = run_broadband_analysis(config(target_voltage_mv=target), library)["estimate"]
        self.assertAlmostEqual(weighted["thickness_scale"], 2.0, places=5)
        self.assertNotAlmostEqual(flat["thickness_scale"], 2.0, places=3)

    def test_partial_coverage_brackets_the_estimate(self):
        result = run_broadband_analysis(config(wavelength_max_nm=3500.0, target_voltage_mv=60.0),
                                        synthetic_library(0.2))
        estimate = result["estimate"]
        self.assertTrue(estimate["partial_coverage"])
        low, high = estimate["thickness_scale_bounds"]
        self.assertLessEqual(low, estimate["thickness_scale"])
        self.assertTrue(high is None or high >= estimate["thickness_scale"])

    def test_voltage_at_or_below_dark_is_refused(self):
        with self.assertRaises(ValueError):
            run_broadband_analysis(config(dark_voltage_mv=10.0, target_voltage_mv=10.0), synthetic_library())

    def test_voltage_above_no_film_reports_zero_and_warns(self):
        result = run_broadband_analysis(config(target_voltage_mv=120.0), synthetic_library(0.2))
        self.assertEqual(result["estimate"]["thickness_scale"], 0.0)
        self.assertEqual(result["estimate"]["status"], "at_or_above_no_film")
        self.assertTrue(any(w.startswith("THICKNESS ESTIMATE") for w in result["warnings"]))

    def test_multiplier_mode_estimates_a_scale_without_layers(self):
        estimate = run_broadband_analysis(config(thickness_mode=REFERENCE_MULTIPLIER, values=[1],
                                                 target_voltage_mv=10 ** -0.5 * 100),
                                          synthetic_library(0.2))["estimate"]
        self.assertAlmostEqual(estimate["thickness_scale"], 2.5, places=6)
        self.assertIsNone(estimate["layer_equivalent"])


class LoaderTests(unittest.TestCase):
    def test_evoh_loader_converts_fraction_transmission(self):
        entry = {"file": "EVOHTransmissionPercentvsWavelength_in_nm_2000nm_to_5000nm.csv",
                 "format": "csv_wavelength_nm", "quantity": "transmittance_fraction", "reference_thickness_um": None}
        material = load_material_spectrum(BASE_DIR, "EVOH", entry)
        import pandas as pd
        frame = pd.read_csv(os.path.join(BASE_DIR, entry["file"]), skipinitialspace=True).sort_values("x")
        np.testing.assert_allclose(material["wavelength_nm"], frame["x"].to_numpy())
        np.testing.assert_allclose(material["absorbance"], -np.log10(frame["y"].to_numpy()))
        self.assertEqual(material["convention"], "base10")
        self.assertIsNone(material["reference_thickness_um"])

    def test_transmittance_units_are_checked(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "t.csv")
            with open(path, "w") as handle:
                handle.write("x,y\n2000,90\n3000,50\n")
            entry = {"file": "t.csv", "format": "csv_wavelength_nm", "reference_thickness_um": 30}
            with self.assertRaises(ValueError):
                load_material_spectrum(tmp, "bad", dict(entry, quantity="transmittance_fraction"))
            percent = load_material_spectrum(tmp, "ok", dict(entry, quantity="transmittance_percent"))
            np.testing.assert_allclose(percent["absorbance"], -np.log10([0.9, 0.5]))
            self.assertEqual(percent["reference_thickness_um"], 30.0)

    def test_bundled_library_has_unknown_reference_thickness(self):
        library = load_material_library(BASE_DIR, include_discovered=False)
        self.assertEqual(set(library), {"Nylon", "PE", "EVOH"})
        self.assertTrue(all(m["reference_thickness_um"] is None for m in library.values()))

    def test_curve_csv_units_and_validation(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "r.csv")
            with open(path, "w") as handle:
                handle.write("wavelength (um),responsivity\n3,0.5\n2,0.2\n")
            curve = load_spectral_curve_csv(path, "responsivity")
            np.testing.assert_allclose(curve["wavelength_nm"], [2000.0, 3000.0])
            np.testing.assert_allclose(curve["values"], [0.2, 0.5])
            with open(path, "w") as handle:
                handle.write("wavelength_nm,responsivity\n2000,-0.1\n3000,0.5\n")
            with self.assertRaises(ValueError):
                load_spectral_curve_csv(path, "responsivity")
            with open(path, "w") as handle:
                handle.write("wavelength_nm,T\n2000,80\n3000,50\n")
            with self.assertRaises(ValueError):
                load_spectral_curve_csv(path, "transmission")

    def test_hpir104_preset_from_component_database(self):
        presets = load_source_presets(BASE_DIR)
        hpir = [p for p in presets.values() if p["name"].startswith("HPIR104")]
        self.assertEqual(len(hpir), 1)
        self.assertAlmostEqual(hpir[0]["temperature_k"], 903.15)


class ExportTests(unittest.TestCase):
    def test_excel_export_matches_pipeline(self):
        from openpyxl import load_workbook
        result = run_broadband_analysis(config(wavelength_max_nm=3500.0, dark_voltage_mv=5.0,
                                               measured_voltages_by_layer={1: 60.0, 3: 30.0}, inspect_value=3),
                                        synthetic_library(0.2))
        with tempfile.TemporaryDirectory() as tmp:
            path = export_broadband_workbook(os.path.join(tmp, "out.xlsx"), result, [])
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["Summary", "Results", "Spectral Calculation", "Layer Predictions",
                                             "Experimental Data", "Charts"])
            results = wb["Results"]
            headers = [cell.value for cell in results[1]]
            predicted_col = headers.index("Predicted Voltage (mV)") + 1
            for i, row in enumerate(result["rows"], start=2):
                self.assertAlmostEqual(results.cell(row=i, column=predicted_col).value, row["predicted_voltage_mv"], places=9)
            self.assertIn("PARTIAL SPECTRAL COVERAGE", str(wb["Summary"]["A4"].value))
            self.assertEqual(wb["Spectral Calculation"].max_row, len(result["grid"]["wavelength_nm"]) + 1)
            self.assertEqual(len(wb["Charts"]._charts), 7)

    def test_excel_export_includes_the_thickness_estimate(self):
        from openpyxl import load_workbook
        result = run_broadband_analysis(config(target_voltage_mv=55.0), synthetic_library(0.2))
        with tempfile.TemporaryDirectory() as tmp:
            path = export_broadband_workbook(os.path.join(tmp, "estimate.xlsx"), result, [])
            summary = load_workbook(path)["Summary"]
            labels = {row[0].value: row[1].value for row in summary.iter_rows(min_col=1, max_col=2) if row[0].value}
            self.assertIn("Thickness estimate from a measured voltage", labels)
            self.assertAlmostEqual(labels["Estimated x / x_ref"], result["estimate"]["thickness_scale"], places=9)
            self.assertAlmostEqual(labels["Measured voltage (mV)"], 55.0, places=9)


def _tk_available():
    try:
        import tkinter
        root = tkinter.Tk()
        root.destroy()
        return True
    except Exception:
        return False


@unittest.skipUnless(_tk_available(), "Tk display not available")
class UiConsistencyTests(unittest.TestCase):
    """The simplified UI must still drive the full pipeline and agree with it."""

    def setUp(self):
        from tkinter import messagebox
        import customtkinter as ctk
        import BroadbandUI

        self.errors = []
        self._original = messagebox.showerror
        messagebox.showerror = lambda title, message, **kw: self.errors.append(message)
        self.addCleanup(self._restore)
        self.root = ctk.CTk()
        self.root.withdraw()
        self.addCleanup(self.root.destroy)
        self.ui = BroadbandUI
        self.window = BroadbandUI.BroadbandAnalysisWindow(self.root)

    def _restore(self):
        from tkinter import messagebox
        messagebox.showerror = self._original

    def test_defaults_are_conservative_and_detector_driven(self):
        window = self.window
        self.assertEqual(window.baseline_var.get(), "Automatic (only if the data needs it)")
        self.assertFalse(window.clamp_var.get() or window.interface_var.get())
        self.assertFalse(window.advanced_visible)          # advanced settings start hidden
        self.assertEqual(window.mode_var.get(), self.ui.MODE_PREDICT)
        self.assertEqual(window.v0_var.get(), "498.0")     # first detector's V0 applied automatically
        window.detector_var.set(window._detector_labels()[1])
        window._detector_changed()
        self.assertEqual(window.v0_var.get(), "73.0")      # switching detector loads its defaults
        self.assertEqual(self.errors, [])

    def test_predict_workflow_matches_core_and_export(self):
        from openpyxl import load_workbook
        window = self.window
        window.material_var.set("Nylon")
        window._material_changed()
        window.value_var.set("3")
        window.update_analysis()
        self.assertEqual(self.errors, [])
        self.assertIsNone(window.result["estimate"])
        row = window.result["rows"][window.result["inspected_index"]]
        self.assertEqual(row["layer_count"], 3)
        self.assertIn("mV predicted", window.headline_label.cget("text"))
        self.assertIn("Nylon", window.subline_label.cget("text"))
        self.assertIn("measured 215.0 mV", window.subline_label.cget("text"))
        self.assertIn("T_eff", window.explain_label.cget("text"))  # the calculation is explained in place
        core = run_broadband_analysis(window.build_config(), window.materials)
        for ui_row, core_row in zip(window.result["rows"], core["rows"]):
            self.assertEqual(ui_row["predicted_voltage_mv"], core_row["predicted_voltage_mv"])
        with tempfile.TemporaryDirectory() as tmp:
            path = export_broadband_workbook(os.path.join(tmp, "ui.xlsx"), window.result, window.experimental_rows())
            sheet = load_workbook(path)["Results"]
            headers = [cell.value for cell in sheet[1]]
            col = headers.index("Predicted Voltage (mV)") + 1
            for i, core_row in enumerate(core["rows"], start=2):
                self.assertAlmostEqual(sheet.cell(row=i, column=col).value, core_row["predicted_voltage_mv"], places=9)

    def test_pe_is_baseline_corrected_automatically_and_reported(self):
        window = self.window
        window.material_var.set("PE")
        window._material_changed()
        window.value_var.set("1")
        window.update_analysis()
        self.assertEqual(self.errors, [])
        self.assertTrue(window.result["corrections"]["baseline_automatic"])
        voltages = [row["predicted_voltage_mv"] for row in window.result["rows"]]
        self.assertTrue(all(b < a for a, b in zip(voltages, voltages[1:])), voltages)
        self.assertIn("Baseline corrected automatically", window.status_label.cget("text"))
        self.assertIn("baseline", window.explain_label.cget("text").lower())

    def test_estimate_workflow_inverts_the_prediction_shown_by_the_ui(self):
        window = self.window
        window.material_var.set("Nylon")
        window._material_changed()
        window.value_var.set("2")
        window.update_analysis()
        predicted = window.result["rows"][window.result["inspected_index"]]["predicted_voltage_mv"]

        window.mode_var.set(self.ui.MODE_ESTIMATE)
        window._mode_changed()
        self.assertIn("Measured detector voltage", window.value_label.cget("text"))
        window.value_var.set(f"{predicted:.10g}")
        window.update_analysis()
        self.assertEqual(self.errors, [])
        estimate = window.result["estimate"]
        self.assertAlmostEqual(estimate["layer_equivalent"], 2.0, places=5)
        self.assertIn("2.00 layers", window.headline_label.cget("text"))
        self.assertIn("from ", window.subline_label.cget("text"))
        self.assertAlmostEqual(window._estimate_x(), 2.0, places=5)

    def test_estimate_sweep_covers_a_thick_film(self):
        window = self.window
        window.material_var.set("Nylon")
        window._material_changed()
        window.clamp_var.set(True)  # monotonic model, so a thick film has a unique solution
        window.value_var.set("7")
        window.update_analysis()
        target = window.result["rows"][window.result["inspected_index"]]["predicted_voltage_mv"]

        window.mode_var.set(self.ui.MODE_ESTIMATE)
        window._mode_changed()
        window.value_var.set(f"{target:.10g}")
        window.update_analysis()
        self.assertEqual(self.errors, [])
        layers = window.result["estimate"]["layer_equivalent"]
        self.assertAlmostEqual(layers, 7.0, places=4)
        # The graph sweep is re-run so it reaches the estimated thickness.
        self.assertGreaterEqual(window.result["rows"][-1]["layer_count"], math.ceil(layers))

    def test_unreachable_voltage_is_reported_not_invented(self):
        window = self.window
        window.material_var.set("Nylon")
        window._material_changed()
        window.mode_var.set(self.ui.MODE_ESTIMATE)
        window._mode_changed()
        window.value_var.set("1")  # far darker than this broadband model can ever go
        window.update_analysis()
        self.assertEqual(self.errors, [])
        self.assertEqual(window.result["estimate"]["status"], "exceeds_max_scale")
        self.assertIn("No thickness fits", window.headline_label.cget("text"))
        self.assertIn("LOWER limit", window.subline_label.cget("text"))

    def test_every_graph_draws_in_both_modes(self):
        window = self.window
        for mode in (self.ui.MODE_PREDICT, self.ui.MODE_ESTIMATE):
            window.mode_var.set(mode)
            window._mode_changed()
            window.update_analysis()
            for graph in self.ui.GRAPH_TYPES:
                window.graph_var.set(graph)
                window.draw_graph()
        self.assertEqual(self.errors, [])

    def test_fractional_layer_count_is_refused_with_guidance(self):
        window = self.window
        window.value_var.set("2.5")
        window.update_analysis()
        self.assertTrue(any("whole numbers" in message for message in self.errors))

    def test_advanced_and_details_panes_toggle(self):
        window = self.window
        window.toggle_advanced()
        self.assertTrue(window.advanced_visible)
        window.toggle_advanced()
        self.assertFalse(window.advanced_visible)
        window.toggle_details()
        self.assertTrue(window.details_visible)
        self.assertIn("WARNING:", window.results_text.get("1.0", "end"))  # full wording lives here
        window.toggle_explanation()
        self.assertTrue(window.explanation_visible)


if __name__ == "__main__":
    unittest.main()
