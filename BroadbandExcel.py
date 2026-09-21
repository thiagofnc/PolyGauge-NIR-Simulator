"""Excel export for broadband detector analysis.

Writes values taken directly from ``run_broadband_analysis``; it performs no
integration or spectral physics of its own.
"""

from __future__ import annotations

from pathlib import Path

# Per-thickness spectral columns are written for at most this many rows so a
# 1000-point sweep does not produce an unusably wide sheet.
MAX_SPECTRAL_COLUMNS = 20


def _number(value):
    """openpyxl cannot write inf/nan; show them as blank cells instead."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    try:
        value = float(value)
    except (TypeError, ValueError):
        return value
    return value if value == value and abs(value) != float("inf") else None


def _text(value):
    if value is None:
        return "—"
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, (tuple, list)):
        return ", ".join(f"{v:g}" if isinstance(v, (int, float)) else str(_text(v)) for v in value) or "—"
    return value if isinstance(value, str) else _number(value)


def export_broadband_workbook(path, result, experimental_rows=None):
    try:
        from openpyxl import Workbook
        from openpyxl.chart import BarChart, Reference, ScatterChart, Series
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("Excel export requires openpyxl. Install project requirements first.") from exc

    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        path = path.with_suffix(".xlsx")
    wb = Workbook()
    header_fill = PatternFill("solid", fgColor="1F4E78")
    section_fill = PatternFill("solid", fgColor="D9E2F3")
    alert_fill = PatternFill("solid", fgColor="F8CBAD")
    header_font = Font(color="FFFFFF", bold=True)
    bold = Font(bold=True)
    rows = result["rows"]
    n_rows = len(rows)
    coverage = result["coverage"]
    corrections = result["corrections"]

    def style_header(sheet, row_index):
        for cell in sheet[row_index]:
            if cell.value is not None:
                cell.fill, cell.font = header_fill, header_font
                cell.alignment = Alignment(wrap_text=True, vertical="center")

    def section(sheet, title):
        sheet.append([title])
        cell = sheet.cell(row=sheet.max_row, column=1)
        cell.font, cell.fill = Font(bold=True, size=12), section_fill

    def pairs(sheet, items):
        for label, value in items:
            sheet.append([label, _text(value)])
            sheet.cell(row=sheet.max_row, column=1).font = bold
            sheet.cell(row=sheet.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    # ---------------------------------------------------------------- Summary
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Broadband Spectral Detector Analysis"])
    ws["A1"].font = Font(size=16, bold=True)
    ws.append(["T(λ,x) = 10^(-A(λ)·x/x_ref);  T_eff = ∫W·T dλ / ∫W dλ · T_interface;  "
               "V = V_dark + (V0 − V_dark)·T_eff"])
    ws.append([])
    if coverage["partial"]:
        ws.append([f"PARTIAL SPECTRAL COVERAGE: material spectrum covers {coverage['fraction']:.1%} of the "
                   "detector/source weight. Predictions and errors are NOT fully valid."])
        ws.cell(row=ws.max_row, column=1).fill = alert_fill
        ws.cell(row=ws.max_row, column=1).font = bold
    for warning in result["warnings"]:
        if not warning.startswith("PARTIAL"):
            ws.append([f"Warning: {warning}"])
    ws.append([])

    material, thickness, weighting, voltages = result["material"], result["thickness"], result["weighting"], result["voltages"]
    section(ws, "Material & thickness")
    pairs(ws, [("Material", material["name"]), ("Spectrum file", material["source"]),
               ("Absorbance convention used", material["convention"]),
               ("Reference thickness x_ref (µm)", material["reference_thickness_um"] or "UNKNOWN"),
               ("Thickness mode", thickness["description"]),
               ("Absolute thickness in µm reported", thickness["reports_um"]),
               ("Material note", material["note"])])
    ws.append([])
    section(ws, "Detector, source & voltages")
    pairs(ws, [("Detector dataset", weighting["detector_name"]),
               ("Detector weighting", weighting["detector_description"]),
               ("Source weighting", weighting["source_description"]),
               ("Optical terms", weighting["optical_terms"] or "none"),
               ("Wavelength window (nm)", weighting["window_nm"]),
               ("Weight domain (nm)", weighting["weight_domain_nm"]),
               ("No-film voltage V0 (mV)", voltages["no_film_voltage_mv"]),
               ("Dark voltage V_dark (mV)", voltages["dark_voltage_mv"])])
    ws.append([])
    section(ws, "Spectral coverage")
    pairs(ws, [("Coverage fraction", coverage["fraction"]), ("Partial coverage", coverage["partial"]),
               ("Material covers (nm)", coverage["covered_range_nm"]),
               ("Negative absorbance points in active region",
                f"{result['negative_absorbance']['points']} of {result['negative_absorbance']['total_points']} "
                f"({result['negative_absorbance']['percent']:.1f}%)")])
    ws.append([])
    section(ws, "Optional corrections")
    pairs(ws, [("Baseline subtraction", corrections["baseline_enabled"]),
               ("Baseline setting", corrections["baseline_mode"]),
               ("Baseline decision", corrections["baseline_reason"]),
               ("Baseline offset subtracted (A)", corrections["baseline_offset"]),
               ("Negative clamp", corrections["clamp_enabled"]),
               ("Points clamped", f"{corrections['clamped_points']} ({corrections['clamped_percent']:.1f}%)"),
               ("Interface correction", corrections["interface_enabled"]),
               ("Refractive index n", corrections["refractive_index"]),
               ("Interface T per layer", corrections["interface_per_layer"]),
               ("Interface assumptions", corrections["interface_assumptions"])])
    ws.append([])

    estimate = result.get("estimate")
    if estimate:
        section(ws, "Thickness estimate from a measured voltage")
        pairs(ws, [("Measured voltage (mV)", estimate["target_voltage_mv"]),
                   ("Measured effective transmission", estimate["target_transmission"]),
                   ("Estimated x / x_ref", estimate["thickness_scale"]),
                   ("Estimated layers", estimate["layer_equivalent"]),
                   ("Estimated thickness x (µm)", estimate["thickness_um"] or "UNKNOWN (x_ref not recorded)"),
                   ("x / x_ref bounds from spectral coverage", estimate["thickness_scale_bounds"]),
                   ("Thickness bounds (µm)", estimate["thickness_um_bounds"]),
                   ("Solver", f"bisection of the same forward model, {estimate['iterations']} iterations, "
                              f"converged={estimate['converged']}, status={estimate['status']}"),
                   ("Back-predicted voltage (mV)", estimate["back_predicted_voltage_mv"]),
                   ("Residual vs measured (mV)", estimate["residual_mv"]),
                   ("Note", estimate["status_note"] or "Solved inside the bracketed range.")])
        ws.append([])

    section(ws, "Model agreement metrics")
    ws.append(["Metric", "Raw model", "Corrected model"])
    style_header(ws, ws.max_row)
    for label, key in [("n compared", "n"), ("MAE (mV)", "mae_mv"), ("RMSE (mV)", "rmse_mv"),
                       ("MAPE (%)", "mape_percent"), ("Agreement R² (1:1 line)", "r_squared")]:
        ws.append([label, _number(result["raw_metrics"][key]), _number(result["metrics"][key])])
    if coverage["partial"]:
        ws.append(["Metrics are computed under PARTIAL SPECTRAL COVERAGE and are not fully valid."])
        ws.cell(row=ws.max_row, column=1).fill = alert_fill

    # ---------------------------------------------------------------- Results
    res = wb.create_sheet("Results")
    res.append(["Layer Count", "Thickness x (µm)", "x / x_ref", "Measured Voltage (mV)", "Raw Predicted Voltage (mV)",
                "Predicted Voltage (mV)", "Predicted V lower bound (mV)", "Predicted V upper bound (mV)",
                "Measured Transmission", "Raw Predicted Transmission", "Predicted Transmission",
                "Measured ln(1/T)", "Predicted ln(1/T)", "Error (mV)", "Absolute Error (mV)", "Percent Error (%)",
                "Absolute Percent Error (%)", "Raw Error (mV)", "Raw Percent Error (%)", "Coverage Fraction",
                "Partial Coverage"])
    style_header(res, 1)
    for row in rows:
        res.append([_number(row["layer_count"]), _number(row["thickness_um"]), _number(row["thickness_scale"]),
                    _number(row["measured_voltage_mv"]), _number(row["raw_predicted_voltage_mv"]),
                    _number(row["predicted_voltage_mv"]), _number(row["predicted_voltage_bounds_mv"][0]),
                    _number(row["predicted_voltage_bounds_mv"][1]), _number(row["measured_transmission"]),
                    _number(row["raw_effective_transmission"]), _number(row["predicted_transmission"]),
                    _number(row["measured_attenuation_natural"]), _number(row["predicted_attenuation_natural"]),
                    _number(row["error_mv"]), _number(row["absolute_error_mv"]), _number(row["percent_error"]),
                    _number(row["absolute_percent_error"]), _number(row["raw_error_mv"]),
                    _number(row["raw_percent_error"]), _number(row["coverage_fraction"]),
                    _number(row["partial_coverage"])])
        if row["partial_coverage"]:
            for cell in res[res.max_row]:
                cell.fill = alert_fill
    first, last = 2, n_rows + 1

    # --------------------------------------------------- Spectral calculation
    grid = result["grid"]
    wavelength = grid["wavelength_nm"]
    n_points = len(wavelength)
    spectral_indices = list(range(min(n_rows, MAX_SPECTRAL_COLUMNS)))
    spec = wb.create_sheet("Spectral Calculation")
    headers = ["Wavelength (nm)", "Inside Material Coverage", "Raw Absorbance (x_ref)", "Absorbance Used (x_ref)",
               "Source S(λ)", "Detector Responsivity R(λ)", "Optical Terms Product", "Weight W(λ)"]
    for i in spectral_indices:
        label = f"x/x_ref={rows[i]['thickness_scale']:g}"
        headers += [f"T(λ) {label}", f"W·T {label}"]
    spec.append(headers)
    style_header(spec, 1)
    terms_product = grid["optical_terms_product"]
    for p in range(n_points):
        values = [float(wavelength[p]), "Yes" if grid["covered_mask"][p] else "No",
                  _number(grid["raw_absorbance"][p]), _number(grid["absorbance_used"][p]),
                  None if grid["source"] is None else float(grid["source"][p]),
                  None if grid["responsivity"] is None else float(grid["responsivity"][p]),
                  None if terms_product is None else float(terms_product[p]), float(grid["weight"][p])]
        for i in spectral_indices:
            values += [_number(result["spectra"][i]["transmission"][p]),
                       _number(result["spectra"][i]["weighted_transmitted"][p])]
        spec.append(values)

    # ------------------------------------------------------ Layer predictions
    layer = wb.create_sheet("Layer Predictions")
    layer.append(["Layer Count", "Thickness x (µm)", "x / x_ref", "∫W dλ (total)", "∫W dλ (covered)",
                  "∫W·T dλ (covered)", "Spectral Transmission", "Interface Transmission", "Effective Transmission",
                  "Effective Attenuation ln(1/T)", "Effective Absorbance log10(1/T)", "Predicted Voltage (mV)"])
    style_header(layer, 1)
    for row in rows:
        layer.append([_number(row["layer_count"]), _number(row["thickness_um"]), _number(row["thickness_scale"]),
                      _number(row["weight_integral_total"]), _number(row["weight_integral_covered"]),
                      _number(row["weighted_transmitted_integral"]), _number(row["spectral_transmission"]),
                      _number(row["interface_transmission"]), _number(row["predicted_transmission"]),
                      _number(row["predicted_attenuation_natural"]), _number(row["predicted_absorbance_base10"]),
                      _number(row["predicted_voltage_mv"])])

    # ------------------------------------------------------ Experimental data
    exp = wb.create_sheet("Experimental Data")
    exp.append(["Material", "Detector", "Layer Count", "Measured Voltage (mV)", "No-film Voltage (mV)",
                "Dark Voltage (mV)"])
    style_header(exp, 1)
    for row in experimental_rows or []:
        exp.append([row.get("material"), row.get("detector"), row.get("layer_count"), row.get("measured_voltage_mv"),
                    row.get("no_film_voltage_mv"), row.get("dark_voltage_mv")])

    # ----------------------------------------------------------------- Charts
    charts = wb.create_sheet("Charts")
    x_col, x_title = (2, "Thickness x (µm)") if result["thickness"]["reports_um"] else (
        (1, "Layer count") if rows and rows[0]["layer_count"] is not None else (3, "x / x_ref"))

    def scatter(title, x_axis, y_axis):
        chart = ScatterChart()
        chart.title, chart.style, chart.height, chart.width = title, 13, 9, 16
        chart.x_axis.title, chart.y_axis.title = x_axis, y_axis
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        return chart

    def add_series(chart, sheet, xc, yc, r0, r1, title, markers=False, line=True):
        series = Series(Reference(sheet, min_col=yc, min_row=r0, max_row=r1),
                        Reference(sheet, min_col=xc, min_row=r0, max_row=r1), title=title)
        series.smooth = False
        if markers:
            series.marker.symbol, series.marker.size = "circle", 7
        else:
            series.marker.symbol = "none"
        if not line:
            series.graphicalProperties.line.noFill = True
        chart.series.append(series)

    suffix = " [PARTIAL COVERAGE]" if coverage["partial"] else ""
    compare = scatter("Measured vs Predicted Voltage" + suffix, x_title, "Detector voltage (mV)")
    add_series(compare, res, x_col, 5, first, last, "Raw prediction", markers=True)
    if corrections["any_enabled"]:
        add_series(compare, res, x_col, 6, first, last, "Corrected prediction", markers=True)
    add_series(compare, res, x_col, 4, first, last, "Measured", markers=True, line=False)
    charts.add_chart(compare, "A1")

    parity = scatter("Parity: Predicted vs Measured" + suffix, "Measured voltage (mV)", "Predicted voltage (mV)")
    add_series(parity, res, 4, 6, first, last, "Predicted", markers=True, line=False)
    charts.add_chart(parity, "K1")

    attenuation = scatter("Effective Attenuation ln(1/T)" + suffix, x_title, "ln(1/T)")
    add_series(attenuation, res, x_col, 13, first, last, "Predicted", markers=True)
    add_series(attenuation, res, x_col, 12, first, last, "Measured", markers=True, line=False)
    charts.add_chart(attenuation, "A20")

    error = BarChart()
    error.title, error.height, error.width = "Prediction Error (%)" + suffix, 9, 16
    error.x_axis.title, error.y_axis.title = "Row", "Error (%)  + = over-predicts"
    error.x_axis.delete = False
    error.y_axis.delete = False
    error.add_data(Reference(res, min_col=16, min_row=1, max_row=last), titles_from_data=True)
    error.set_categories(Reference(res, min_col=x_col, min_row=first, max_row=last))
    charts.add_chart(error, "K20")

    absorbance = scatter("Reference Absorbance Spectrum", "Wavelength (nm)", "Absorbance")
    add_series(absorbance, spec, 1, 3, 2, n_points + 1, "Raw")
    if corrections["any_enabled"]:
        add_series(absorbance, spec, 1, 4, 2, n_points + 1, "Used")
    charts.add_chart(absorbance, "A39")

    transmission = scatter("Spectral Transmission T(λ)", "Wavelength (nm)", "T")
    for column_index, i in enumerate(spectral_indices):
        add_series(transmission, spec, 1, 9 + 2 * column_index, 2, n_points + 1,
                   f"x/x_ref={rows[i]['thickness_scale']:g}")
    charts.add_chart(transmission, "K39")

    inspected = result["inspected_index"] if result["inspected_index"] in spectral_indices else 0
    contribution = scatter(f"Weighted Contribution (x/x_ref={rows[inspected]['thickness_scale']:g})",
                           "Wavelength (nm)", "Weighted value")
    add_series(contribution, spec, 1, 8, 2, n_points + 1, "W(λ)")
    add_series(contribution, spec, 1, 10 + 2 * inspected, 2, n_points + 1, "W·T")
    charts.add_chart(contribution, "A58")

    # ----------------------------------------------------------------- Layout
    ws.column_dimensions["A"].width = 44
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 18
    for sheet in (res, spec, layer, exp):
        sheet.freeze_panes = "B2"
        for column_index in range(1, sheet.max_column + 1):
            header = sheet.cell(row=1, column=column_index).value
            sheet.column_dimensions[get_column_letter(column_index)].width = max(14, min(30, len(str(header or "")) + 3))
    wb.save(path)
    return str(path)
